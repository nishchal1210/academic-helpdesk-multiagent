import hashlib
import json

import pytest

from helpdesk.corpus import CorpusError, Registry, extract_pages
from helpdesk.retrieval import Retriever, ingest
from helpdesk.grounding import validate_answer
from helpdesk.schemas import ExaminationRow
from helpdesk.tools import NoEvidence
from helpdesk.scope import route_question


def test_spreadsheet_preserves_sheet_row_and_coordinates(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "schedule.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Course slots"
    ws["B3"] = "CSL7650"
    ws["E3"] = "E"
    ws["A1000"].number_format = "0"
    wb.save(path)
    pages = list(extract_pages(path))
    assert len(pages) == 1
    assert "row 3" in pages[0][0] and "Course slots" in pages[0][0]
    assert "B3: CSL7650" in pages[0][1] and "E3: E" in pages[0][1]


def test_ocr_sidecar_is_bound_to_source_hash(tmp_path):
    from pypdf import PdfWriter
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(path)
    sidecar = path.with_suffix(".pdf.extraction.json")
    text = "A revised attendance policy recovered from the scanned original page."
    sidecar.write_text(json.dumps({"source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                                  "pages": {"1": text}}))
    assert list(extract_pages(path)) == [("page 1", text)]
    sidecar.write_text(json.dumps({"source_sha256": "stale", "pages": {"1": text}}))
    with pytest.raises(CorpusError, match="Stale"):
        list(extract_pages(path))


def test_pending_curriculum_is_not_used(settings):
    p = settings.data_dir / "manifest.json"
    data = json.loads(p.read_text())
    data["documents"][1]["status"] = "pending_review"
    data["documents"][1]["effective_date"] = None
    p.write_text(json.dumps(data))
    ingest(settings, lexical_only=True)
    assert all(e.document_id != "demo-degree" for e in Retriever(settings, Registry(settings.data_dir)).records.values())


def test_missing_email_does_not_break_schedule_lookup(tools):
    tools.registry.faculty[0].email = ""
    result = tools.lookup_course("CSL7570")
    assert result["offering_rows"] == 1
    with pytest.raises(NoEvidence, match="email"):
        tools.draft_email("Draft sick leave email for CSL7570 for 2 days")


def test_application_attaches_exact_passage_for_valid_source_id(tools):
    tools.lookup_course("CSL7570")
    evidence = next(iter(tools.ledger.items.values()))
    raw = json.dumps({"status": "answered", "claims": [{"text": "The course uses slot D3.",
            "supports": [{"evidence_id": evidence.id}]}]})
    answer = validate_answer(raw, tools.ledger)
    assert answer.claims[0].supports[0].quote == evidence.text


def test_unknown_elective_type_prevents_count(tools):
    tools.registry.courses[1].course_type = "unknown"
    with pytest.raises(NoEvidence, match="classification"):
        tools.count_electives("Humanities", 400)


def test_department_completeness_does_not_imply_other_departments(tools):
    tools.registry.manifest.complete_course_semesters.clear()
    tools.registry.manifest.complete_department_semesters = {tools.settings.semester: ["Humanities"]}
    assert tools.count_electives("HSS", 400)["count"] == 3
    with pytest.raises(NoEvidence, match="not marked complete"):
        tools.count_electives("Computer Science", 400)


def test_unconfirmed_exam_shift_cannot_be_presented_as_final(tools):
    tools.registry.examinations = [ExaminationRow(course_code="CSL7570", examination="major",
        semester=tools.settings.semester, date="2026-11-22", time="09:00-12:00", status="unconfirmed_change",
        source_document_id="demo-timetable", source_locator="row 4")]
    with pytest.raises(NoEvidence, match="without confirmed approval"):
        tools.lookup_examination("CSL7570", "major")


def test_makeup_major_date_is_not_regular_major_timetable(tools):
    question = "When is the makeup Major examination for CSL7570?"
    assert route_question(question).name == "policy"
    with pytest.raises(NoEvidence, match="does not establish"):
        tools.schedule_for_question(question)


def test_exam_lookup_cites_course_slot_and_timetable(tools):
    tools.registry.examinations = [ExaminationRow(course_code="CSL7570", examination="major",
        semester=tools.settings.semester, date="2026-11-22", time="09:00-12:00", status="tentative",
        source_document_id="demo-timetable", source_locator="row 4")]
    result = tools.lookup_examination("CSL7570", "major")
    assert any("tentative" in e["text"] for e in result["evidence"])
    assert any("slot D3" in e["text"] for e in result["evidence"])


def test_topic_replacement_excludes_historical_document(settings):
    p = settings.data_dir / "manifest.json"
    data = json.loads(p.read_text())
    data["topic_overrides"] = [{"keywords": ["attendance"], "preferred_document_ids": ["demo-attendance"],
                                "excluded_document_ids": ["demo-degree"]}]
    p.write_text(json.dumps(data))
    ingest(settings, lexical_only=True)
    r = Retriever(settings, Registry(settings.data_dir))
    assert all(e.document_id != "demo-degree" for e in r.search("attendance degree"))


def test_sidecar_change_invalidates_index(indexed):
    path = indexed.data_dir / "documents/attendance.md.extraction.json"
    path.write_text(json.dumps({"source_sha256": "anything", "pages": {}}))
    with pytest.raises(CorpusError, match="changed"):
        Retriever(indexed, Registry(indexed.data_dir))
