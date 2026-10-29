"""Validated source registry and page-preserving document extraction."""
import csv
from datetime import date
import hashlib
import json
from pathlib import Path
import re

from .schemas import CourseRow, Document, ExaminationRow, FacultyRow, Manifest


def tokens(text: str) -> list[str]:
    # Preserve course codes and normalize only common inflections.
    words = re.findall(r"[a-z0-9]+", text.lower())
    stop = {"the", "a", "an", "is", "are", "of", "to", "for", "in", "and", "what", "how", "do", "i", "my", "this"}
    return [w[:-1] if len(w) > 4 and w.endswith("s") else w for w in words if w not in stop]


class CorpusError(ValueError):
    pass


class Registry:
    def __init__(self, data_dir: Path):
        self.data_dir = data_dir.resolve()
        path = self.data_dir / "manifest.json"
        if not path.is_file():
            raise CorpusError(f"Missing {path}. Add your source manifest or use --demo.")
        self.manifest = Manifest.model_validate_json(path.read_text())
        self.docs = {d.id: d for d in self.manifest.documents}
        if len(self.docs) != len(self.manifest.documents):
            raise CorpusError("Document IDs must be unique.")
        for d in self.docs.values():
            if d.effective_date and date.fromisoformat(d.effective_date) > date.today():
                raise CorpusError(f"{d.id}: future effective date; do not mark as currently applicable.")
            if not self.path(d).is_file():
                raise CorpusError(f"Missing document: {d.file}")
            for old in d.supersedes:
                if old not in self.docs:
                    raise CorpusError(f"{d.id}: unknown superseded document {old}")
            if d.synthetic and not self.manifest.synthetic:
                raise CorpusError("Synthetic documents cannot be mixed into the real corpus.")
            if self.manifest.synthetic and not d.synthetic:
                raise CorpusError("Every document in a synthetic corpus must be labelled synthetic.")
        self.courses = self._read_csv("courses.csv", CourseRow)
        self.faculty = self._read_csv("faculty.csv", FacultyRow)
        self.examinations = self._read_csv("examinations.csv", ExaminationRow)
        self._validate_tables()

    def path(self, doc: Document) -> Path:
        path = (self.data_dir / doc.file).resolve()
        if not path.is_relative_to(self.data_dir):
            raise CorpusError("Document paths must remain inside the data directory.")
        return path

    def _read_csv(self, name, cls):
        path = self.data_dir / name
        if not path.is_file():
            return []
        with path.open(encoding="utf-8-sig", newline="") as f:
            return [cls.model_validate(row) for row in csv.DictReader(f)]

    def _validate_tables(self):
        ids = [f.instructor_id for f in self.faculty]
        if len(ids) != len(set(ids)):
            raise CorpusError("Faculty instructor_id values must be unique.")
        seen = set()
        for row in [*self.courses, *self.faculty, *self.examinations]:
            if row.source_document_id not in self.docs or not row.source_locator.strip():
                raise CorpusError("Every CSV row needs a registered source document and a source locator.")
        for row in self.faculty:
            if row.email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", row.email):
                raise CorpusError(f"Invalid faculty email for {row.instructor_id}")
        for row in self.courses:
            if not re.fullmatch(r"[A-Z]{2,5}\d{3,5}[A-Z]?", row.course_code):
                raise CorpusError(f"Non-canonical course code: {row.course_code}")
            key = (row.course_code, row.semester, row.section, row.instructor_id)
            if key in seen:
                raise CorpusError(f"Duplicate course row: {key}")
            seen.add(key)
            if row.instructor_id not in ids:
                raise CorpusError(f"{row.course_code}: missing faculty record {row.instructor_id}")
            if row.slot_document_id and row.slot_document_id not in self.docs:
                raise CorpusError(f"{row.course_code}: unknown slot-system source")

    def applicable(self, doc_id, semester="", programme="", cohort="", branch=""):
        d = self.docs[doc_id]
        superseded = {old for doc in self.docs.values() if doc.status == "current" for old in doc.supersedes}
        return (d.status == "current" and d.id not in superseded
                and (not d.semesters or semester in d.semesters)
                and (not d.programmes or programme in d.programmes)
                and (not d.cohorts or cohort in d.cohorts)
                and (not d.branches or branch in d.branches))

    def fingerprint(self):
        h = hashlib.sha256()
        h.update(self.manifest.model_dump_json().encode())
        for doc in self.docs.values():
            with self.path(doc).open("rb") as f:
                for block in iter(lambda: f.read(1024 * 1024), b""):
                    h.update(block)
            sidecar = self.path(doc).with_suffix(self.path(doc).suffix + ".extraction.json")
            if sidecar.exists():
                h.update(sidecar.read_bytes())
        for name in ["courses.csv", "faculty.csv", "examinations.csv"]:
            path = self.data_dir / name
            if path.exists():
                h.update(path.read_bytes())
        return h.hexdigest()


def extract_pages(path: Path):
    suffix = path.suffix.lower()
    sidecar_path = path.with_suffix(path.suffix + ".extraction.json")
    sidecar = {}
    if sidecar_path.exists():
        sidecar = json.loads(sidecar_path.read_text())
        if sidecar.get("source_sha256") != hashlib.sha256(path.read_bytes()).hexdigest():
            raise CorpusError(f"Stale extraction/OCR sidecar for {path.name}; prepare it again.")
        sidecar = sidecar.get("pages", {})
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise CorpusError(f"Encrypted PDF: {path.name}")
        for i, page in enumerate(reader.pages, 1):
            text = sidecar.get(str(i), page.extract_text(extraction_mode="layout") or "")
            if len(text.strip()) < 30:
                raise CorpusError(f"{path.name} page {i}: little or no text. OCR this page, then ingest again.")
            yield f"page {i}", text
    elif suffix == ".docx":
        from docx import Document as WordDocument
        doc = WordDocument(path)
        parts = [p.text for p in doc.paragraphs]
        parts += [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        yield "DOCX text (no stable pagination)", "\n".join(parts)
        for locator, text in sidecar.items():
            yield locator, text
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            for sheet in workbook:
                # Empty formatted rows are not documents; retain actual cell coordinates and sheet names.
                for row_number, row in enumerate(sheet, 1):
                    cells = [f"{c.coordinate}: {c.value}" for c in row if c.value is not None
                             and str(c.value).strip()]
                    if cells:
                        yield f"sheet {sheet.title!r}, row {row_number}", "; ".join(cells)
        finally:
            workbook.close()
    elif suffix in {".txt", ".md"}:
        yield "text", path.read_text(encoding="utf-8-sig")
    else:
        raise CorpusError(f"Unsupported format {suffix}: use PDF, DOCX, XLSX, TXT or MD.")


def chunks(text: str, size=900, overlap=150):
    """Character-bounded chunks with overlap; never join text across PDF pages."""
    text = re.sub(r"[ \t]+", " ", text).strip()
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            split = max(text.rfind("\n", start + size // 2, end), text.rfind(". ", start + size // 2, end))
            if split > start:
                end = split + 1
        part = text[start:end].strip()
        if part:
            yield part
        if end == len(text):
            break
        start = max(start + 1, end - overlap)


def atomic_json(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2))
    temporary.replace(path)
